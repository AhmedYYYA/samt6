"""Build the complete bilingual Word proposal from the shared programme content."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
DATA=json.loads((ROOT/'docs/programme-content.json').read_text())
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    st.element.rPr.rFonts.set(qn('w:cs'),'Noto Sans Arabic')
    st.paragraph_format.space_after=Pt(8)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
for n,size in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    doc.styles[n].font.size=Pt(size);doc.styles[n].font.bold=True
    doc.styles[n].paragraph_format.space_before=Pt(15 if n!='Title' else 0)
    doc.styles[n].paragraph_format.keep_with_next=True
doc.core_properties.title='SAMT Ministry of Defence leadership development proposal'
doc.core_properties.subject='Arabic and English proposal for in-principle endorsement and pilot design'
doc.core_properties.author='Ahmed Younis Yousif Alhammadi'
doc.core_properties.keywords='SAMT, leadership, MOD, proposal'

def fmt(p,ar=False,size=None):
    p.alignment=None
    pp=p._p.get_or_add_pPr()
    bd=OxmlElement('w:bidi');bd.set(qn('w:val'),'1' if ar else '0');pp.append(bd)
    jc=OxmlElement('w:jc');jc.set(qn('w:val'),'start');pp.append(jc)
    for r in p.runs:
        r.font.name='Noto Sans Arabic' if ar else 'Arial'
        rp=r._r.get_or_add_rPr(); rf=rp.rFonts
        if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        for k in ['ascii','hAnsi','cs']:rf.set(qn('w:'+k),'Noto Sans Arabic' if ar else 'Arial')
        rtl=OxmlElement('w:rtl');rtl.set(qn('w:val'),'1' if ar else '0');rp.append(rtl)
        actual=size or (p.style.font.size.pt if p.style.font.size else 10.5)
        r.font.size=Pt(actual)
        sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(int(actual*2)));rp.append(sz)
    return p

def paragraph(text,ar=False,style=None):
    p=doc.add_paragraph(text,style=style);fmt(p,ar)
    return p

def heading(text,ar=False,level=1):
    p=doc.add_heading(text,level=level);fmt(p,ar,15 if level==1 else 12)
    return p

def table(rows,ar=False,widths=None):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    props=t._tbl.tblPr
    if ar:props.append(OxmlElement('w:bidiVisual'))
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        if i==0:t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        else:t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,txt in enumerate(row):
            c=cells[j];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Inches(widths[j])
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if i==0 else ('F5F5F5' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(sh)
            p=c.paragraphs[0];p.add_run(txt);fmt(p,ar,9.3)
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.12
            if i==0:
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def hyperlink(p,text,url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'333333');rp.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);run.append(rp);tx=OxmlElement('w:t');tx.text=text;run.append(tx);link.append(run);p._p.append(link)

for lang in ['ar','en']:
    ar=lang=='ar';text=lambda x:x[lang] if isinstance(x,dict) else x
    if not ar:doc.add_page_break()
    paragraph('مقترح اعتماد مبادرة سمت لإعداد قيادات وزارة الدفاع' if ar else 'SAMT proposal for Ministry of Defence leadership development',ar,'Title')
    paragraph('مقدم المبادرة أحمد يونس يوسف الحمادي\n25 سبتمبر 2026  |  الإصدار 1.0\nالنسخة العربية تتبعها النسخة الإنجليزية المكافئة' if ar else 'Proposed by Ahmed Younis Yousif Alhammadi\n25 September 2026  |  Version 1.0\nEnglish equivalent of the preceding Arabic proposal',ar)
    for s in DATA['proposal']:
        heading(text(s['title']),ar)
        for p in s.get('paragraphs',[]):paragraph(text(p),ar)
        if 'rows'in s:
            widths=[1.25,1.65,4] if s['id']=='gates' else [1.45,3.85,1.6]
            table([[text(c) for c in r] for r in s['rows']],ar,widths)
    doc.add_page_break()
    heading('ملحق المحطات والجهات المقترحة' if ar else 'Appendix on stations and proposed hosts',ar)
    for s in DATA['stations']:
        heading(text(s['country']),ar,2)
        paragraph(('الأسابيع ' if ar else 'Weeks ')+f"{s['start']}–{s['end']}  |  "+text(s['city']),ar)
        paragraph(text(s['description']),ar)
        paragraph(('الجهات والمسارات المقترحة: ' if ar else 'Proposed hosts and tracks: ')+'؛ '.join(text(h) for h in s['hosts']),ar)
        paragraph(('المخرج: ' if ar else 'Output: ')+text(s['output']),ar)
        paragraph(text(s['note']),ar)
    heading('ملحق خطة الأسابيع' if ar else 'Appendix on the weekly plan',ar)
    table(([['الأسبوع','الموضوع','المخرج المتوقع']] if ar else [['Week','Theme','Expected output']])+[[str(i+1),text(h),text(p)] for i,(h,p) in enumerate(DATA['weeks'])],ar,[.55,2.1,4.25])
    heading('ملحق الجدارات والأدلة' if ar else 'Appendix on competencies and evidence',ar)
    paragraph('الأوزان التالية مقترحة للمعايرة ومجموعها 100 بالمئة. يراجع المقيم الأدلة والسلوك قبل إصدار الحكم.' if ar else 'The following weights are proposed for calibration and total 100 percent. Assessors review evidence and behaviour before forming a judgement.',ar)
    for d in DATA['domains']:
        heading(text(d['name']),ar,2)
        paragraph(('الوزن المقترح ' if ar else 'Proposed weight ')+str(d['weight'])+'%',ar).paragraph_format.keep_with_next=True
        paragraph(text(d['description'])+' '+text(d['evidence']),ar)
    heading('المراجع وحدود الاستدلال' if ar else 'References and limits of evidence',ar)
    paragraph('تسند المراجع عناصر التصميم أو الاختصاص المعلن للجهات. ولا تثبت اعتماد المبادرة أو توافر الاستضافة. تاريخ مراجعة الروابط 25 سبتمبر 2026.' if ar else 'References support elements of the design or the stated institutional remit. They do not establish initiative approval or hosting availability. Links reviewed on 25 September 2026.',ar)
    for s in DATA['sources']:
        p=paragraph('['+s['id']+'] '+text(s['title'])+' — '+s['publisher'],ar);p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(3)
        p=paragraph(text(s['use']),ar);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
        p=doc.add_paragraph();hyperlink(p,s['url'],s['url']);fmt(p,False,8)
        p.paragraph_format.space_after=Pt(5)

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run('SAMT   |   ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for element in [doc.styles.element,doc._element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(ROOT/'docs/SAMT_Approval_Proposal.docx')
print(ROOT/'docs/SAMT_Approval_Proposal.docx')
